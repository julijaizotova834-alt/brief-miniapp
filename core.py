"""
core.py — Skrepka AI (українська версія)
Загальна логіка для бота та міні-додатку.
"""

import os, re, logging, tempfile, shutil
from io import BytesIO
from datetime import datetime

import psycopg2
from groq import Groq
import google.generativeai as genai
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

logger = logging.getLogger(__name__)

def _groq_client():
    key = os.environ.get('GROQ_API_KEY')
    if not key:
        raise RuntimeError('GROQ_API_KEY is not configured')
    return Groq(api_key=key)

def _gemini_model():
    key = os.environ.get('GEMINI_API_KEY')
    if not key:
        raise RuntimeError('GEMINI_API_KEY is not configured')
    genai.configure(api_key=key)
    return genai.GenerativeModel('gemini-2.5-flash')


BOT_NAME = 'skrepka'
FREE_LIMIT = 7
GROQ_SIZE_LIMIT_MB = 24
PRICE_SUB = 750
SUB_DAYS = 30
AUDIT_PACKS = [
    {'size': 10, 'price': 150}, {'size': 50, 'price': 550}, {'size': 100, 'price': 900},
]
FREE_USERNAMES = {'black_physicist', 'juli_aleksandrova', 'aliia85'}

PROMPTS = {
    'narada': """Проаналізуй транскрипцію зустрічі та підготуй короткий звіт українською мовою. Структура: 1. Резюме зустрічі (3-5 речень — суть, мета, результат) 2. Основні теми розмови (короткий перелік) 3. Ключові рішення (короткий перелік) 4. Домовленості сторін (короткий перелік) 5. Завдання — кожне в одному рядку за шаблоном: • [Відповідальний] — [Завдання] | Термін: [термін] | [Коментар якщо є] 6. Важливі цифри, суми та терміни 7. Відкриті питання 8. Наступні кроки. Правила: не вигадуй інформацію, якої немає в записі; при помилках розпізнавання — логічно відновлюй зміст; НЕ використовуй таблиці в жодному вигляді; лише заголовки, списки та абзаци; діловий стиль без зайвої «води».""",

    'audit': """Дій як РВП (керівник відділу продажів) та бізнес-тренер з переговорів. Проаналізуй транскрипцію розмови менеджера з клієнтом (або замовником) і сформуй повний розбір дзвінка українською мовою. ПРАВИЛА ФОРМАТУВАННЯ — виконуй суворо: - НЕ використовуй таблиці в жодному вигляді - Кожен пункт — це один короткий абзац або один рядок списку - Лише заголовки, марковані списки та короткі абзаци - Діловий, але прямий стиль — без «води» та повторів ПРАВИЛА ЗМІСТУ: - Не вигадуй інформацію, якої немає в транскрипції - Якщо даних немає — «Не виявлено» або «Не прозвучало в розмові» - При помилках розпізнавання — логічно відновлюй зміст без зміни суті - Будь чесним в оцінці менеджера — не прикрашай і не пом'якшуй СТРУКТУРА: 1. Короткий підсумок Один абзац 3-4 речення: хто з ким спілкувався, мета дзвінка, чим закінчилася розмова, загальне враження. 2. Профіль клієнта На основі того, що клієнт сказав і як себе поводив: • Тип клієнта (гарячий / теплий / холодний) • Рівень обізнаності про продукт/послугу • Ключовий запит — що саме клієнт хоче вирішити • Прихована потреба — що стоїть за запитом, якщо можна визначити • Болі клієнта — які проблеми, страхи чи незручності озвучив • Критерії прийняття рішення — за якими параметрами обирає (ціна, терміни, довіра, досвід) • Бюджет та очікування щодо ціни — якщо прозвучало 3. Заперечення та сумніви клієнта Кожне заперечення у форматі: • [Заперечення] — [як менеджер відпрацював чи не відпрацював] Якщо заперечень не було — «Заперечень не прозвучало.» 4. Комерційні умови Якщо обговорювалися ціни, терміни, обсяги — коротким переліком: • [Що]: [сума / термін / умова] Якщо не обговорювалися — «Комерційні умови не порушувалися.» 5. Оцінка менеджера Оціни роботу менеджера за кожним критерієм — коротко, конкретно, з прикладами з розмови: • Встановлення контакту — привітався, представився, задав тон розмові чи ні • Виявлення потреби — ставив відкриті питання чи одразу почав продавати • Активне слухання — чув клієнта, уточнював, перефразовував чи перебивав і гнув своє • Презентація рішення — подав продукт через вигоди клієнта чи через характеристики • Робота із запереченнями — відпрацював, проігнорував, сперечався чи погодився і здався • Закриття / наступний крок — запропонував конкретну дію чи відпустив клієнта без домовленості • Тон і манера — впевненість, дружелюбність, тиск, невпевненість, метушня 6. Помилки менеджера Конкретний перелік з прикладами — що менеджер зробив не так: • [Помилка] — [що було сказано чи зроблено] — [як треба було] Якщо критичних помилок немає — зазначай точки зростання: що можна було зробити краще. 7. Сильні сторони менеджера Що менеджер зробив добре — конкретно, з прикладами з розмови. 8. Етап угоди та прогноз Одним абзацом: • На якому етапі воронки клієнт (первинний контакт / виявлення потреби / презентація / обговорення умов / прийняття рішення / відмова) • Ймовірність угоди (висока / середня / низька) — з аргументом чому • Що може завадити закриттю 9. Завдання та наступні кроки Кожне завдання — ОДИН рядок: • [Відповідальний] — [Завдання] | Термін: [термін] Якщо відповідальний невідомий — «Уточнити». Якщо термін невідомий — «Не визначено». 10. Рекомендації менеджеру 3-5 конкретних рекомендацій що робити з цим клієнтом далі та що покращити в техніці продажів для подібних дзвінків.""",

    'client_summary': """Дій як професійний асистент керівника. Проаналізуй транскрипцію зустрічі або дзвінка з клієнтом (замовником, партнером) і сформуй резюме зустрічі українською мовою, яке можна надіслати клієнту одразу після розмови. ПРАВИЛА ФОРМАТУВАННЯ — виконуй суворо: - НЕ використовуй таблиці в жодному вигляді - Лише заголовки, короткі списки та абзаци - Ввічливий, професійний діловий стиль - Документ має виглядати як готовий лист або повідомлення для надсилання клієнту - НЕ додавай внутрішні позначки, оцінки чи коментарі — клієнт це побачить ПРАВИЛА ЗМІСТУ: - Не вигадуй інформацію, якої немає в транскрипції - Якщо щось не обговорювалося — не згадуй цей пункт взагалі - При помилках розпізнавання — логічно відновлюй зміст - НЕ включай внутрішні завдання команди — лише те, що стосується клієнта - НЕ включай ціни, бюджети та фінансові умови, якщо вони не були остаточно узгоджені обома сторонами — чорнові обговорення цін не виносити в резюме - Тон: шанобливий, конкретний, без підлабузництва та канцеляриту СТРУКТУРА: Звернення «Доброго дня, [ім'я клієнта якщо прозвучало / колеги]!» і одне речення: «Надсилаю резюме нашої зустрічі від [дата якщо відома].» 1. Що обговорили Компактний перелік ключових тем розмови — 3-7 пунктів, кожен в один рядок: • [Тема]: [суть в одному реченні] Без внутрішніх деталей — лише те, що важливо для клієнта. 2. Домовленості Чіткий перелік того, про що домовилися обидві сторони: • [Що узгоджено] — [хто робить / що відбувається далі] Лише підтверджені обома сторонами рішення. 3. Наступні кроки Перелік конкретних дій з нашого боку та з боку клієнта: • З нашого боку: [дія] — [термін якщо обговорювався] • З вашого боку: [дія] — [термін якщо обговорювався] Якщо терміни не обговорювалися — не вигадуй їх. 4. Відкриті питання Якщо залишилися питання, що потребують уточнення або додаткового обговорення — короткий перелік. Якщо все вирішено — цей розділ не включай. Завершення Одне речення: «Якщо я щось пропустив або ви хочете доповнити — буду радий зворотному зв'язку.» і підпис: «З повагою, [місце для імені]» НЕ додавай нічого після завершення. Документ має бути готовий до копіювання та надсилання клієнту.""",

    'ideas': """Дій як помічник керівника та стратегічний консультант. Проаналізуй голосові нотатки або мозковий штурм і перетвори їх на структурований план дій українською мовою. Правила: не транскрибуй текст дослівно — структуруй і відбирай головне; не вигадуй цифри, бюджети чи терміни, яких не було; якщо даних немає — «Не визначено»; НЕ використовуй таблиці в жодному вигляді — лише заголовки, марковані списки та абзаци; управлінський стиль без «води»; відповідай компактно і без повторів. Структура: 1. Головна суть і мета (3-4 речення) 2. Ключові ідеї та їхня цінність (перелік) 3. Стратегічні напрямки та перспективи 4. Операційні завдання — кожне в одному рядку: • [Завдання] | Відповідальний: [хто] | Пріоритет: [високий/середній/низький] | Термін: [коли] 5. Термінові питання та критичні блокери 6. Що делегувати і кому (ролі або спеціалісти) 7. Ризики — фінансові, юридичні, технічні або операційні (коротко) 8. Ресурси — люди, інструменти, бюджет 9. Наступні кроки — план на 24 год / 7 днів / 30 днів 10. Рекомендація — що найперспективніше і що перевірити в першу чергу. Якщо в записі є ідея продукту або бізнесу — додатково: цільова аудиторія, основна цінність, мінімальна версія, монетизація, ключові функції, складність запуску.""",
}

DOC_NAMES = {
    'narada': 'Нарада', 'audit': 'Аудит дзвінка менеджера',
    'client_summary': 'Резюме для клієнта', 'ideas': 'Структуровані ідеї',
    'transcript': 'Транскрипція',
}

def get_db(): return psycopg2.connect(os.environ.get('DATABASE_URL'))

def init_db():
    conn = get_db(); cur = conn.cursor()
    cur.execute('''CREATE TABLE IF NOT EXISTS users (
        user_id BIGINT, bot_name TEXT, username TEXT, full_name TEXT,
        registered_at TIMESTAMP DEFAULT NOW(), usage_count INTEGER DEFAULT 0,
        is_subscribed BOOLEAN DEFAULT FALSE, subscription_until TIMESTAMP,
        audit_credits INTEGER DEFAULT 0, last_used_at TIMESTAMP,
        PRIMARY KEY (user_id, bot_name))''')
    cur.execute('''CREATE TABLE IF NOT EXISTS payments (
        id SERIAL PRIMARY KEY, user_id BIGINT, bot_name TEXT,
        telegram_payment_charge_id TEXT UNIQUE, stars_amount INTEGER,
        payload TEXT, paid_at TIMESTAMP DEFAULT NOW())''')
    for col, ct, df in [('audit_credits','INTEGER','0')]:
        try: cur.execute(f"ALTER TABLE users ADD COLUMN {col} {ct} DEFAULT {df}")
        except: conn.rollback()
    try: cur.execute("ALTER TABLE payments ADD COLUMN payload TEXT")
    except: conn.rollback()
    try: cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_pay_charge ON payments(telegram_payment_charge_id)")
    except: conn.rollback()
    conn.commit(); cur.close(); conn.close()

def register_user_if_new(uid, uname, fname):
    conn = get_db(); cur = conn.cursor()
    cur.execute('''INSERT INTO users (user_id,bot_name,username,full_name,registered_at,usage_count,audit_credits,last_used_at)
        VALUES (%s,%s,%s,%s,NOW(),0,0,NOW()) ON CONFLICT (user_id,bot_name) DO UPDATE
        SET username=EXCLUDED.username, full_name=EXCLUDED.full_name, last_used_at=NOW()''', (uid, BOT_NAME, uname, fname))
    conn.commit(); cur.close(); conn.close()

def increment_usage(uid):
    conn = get_db(); cur = conn.cursor()
    cur.execute('UPDATE users SET usage_count=usage_count+1, last_used_at=NOW() WHERE user_id=%s AND bot_name=%s', (uid, BOT_NAME))
    conn.commit(); cur.close(); conn.close()

def get_user_info(uid):
    conn = get_db(); cur = conn.cursor()
    cur.execute('SELECT usage_count,is_subscribed,subscription_until,audit_credits FROM users WHERE user_id=%s AND bot_name=%s', (uid, BOT_NAME))
    r = cur.fetchone(); cur.close(); conn.close()
    if not r: return None
    return {'usage': r[0], 'sub': r[1], 'until': r[2], 'credits': r[3] or 0}

def has_active_sub(info): return info and info['sub'] and info['until'] and info['until'] > datetime.now()

def check_access(uid, uname):
    if uname and uname.lower() in FREE_USERNAMES: return True
    info = get_user_info(uid)
    if not info: return True
    if has_active_sub(info): return True
    return info['usage'] < FREE_LIMIT

def check_audit_access(uid, uname):
    """Повертає free (привілейований), free_ok (безкоштовна квота), credit (платний кредит), або need_credits."""
    if uname and uname.strip().lstrip('@').lower() in FREE_USERNAMES:
        return 'free'
    info = get_user_info(uid)
    if not info:
        return 'free_ok'
    if has_active_sub(info):
        return 'credit' if info['credits'] > 0 else 'need_credits'
    if info['credits'] > 0:
        return 'credit'
    if info['usage'] < FREE_LIMIT:
        return 'free_ok'
    return 'need_credits'


def reserve_free_usage(uid, uname=''):
    """Атомарно резервує одну безкоштовну спробу."""
    if uname and uname.strip().lstrip('@').lower() in FREE_USERNAMES:
        return True
    conn = get_db()
    try:
        with conn.cursor() as cur:
            cur.execute("""UPDATE users SET usage_count=usage_count+1, last_used_at=NOW()
                WHERE user_id=%s AND bot_name=%s AND usage_count < %s
                AND NOT (is_subscribed=TRUE AND subscription_until>NOW())
                RETURNING usage_count""", (uid, BOT_NAME, FREE_LIMIT))
            if cur.fetchone():
                conn.commit(); return True
            cur.execute("""SELECT 1 FROM users WHERE user_id=%s AND bot_name=%s
                AND is_subscribed=TRUE AND subscription_until>NOW()""", (uid, BOT_NAME))
            subscribed = cur.fetchone() is not None
            conn.commit()
            return subscribed
    except Exception:
        conn.rollback(); raise
    finally:
        conn.close()


def release_free_usage(uid):
    """Відкат раніше зарахованої безкоштовної спроби після помилки обробки/доставки."""
    conn = get_db()
    try:
        with conn.cursor() as cur:
            cur.execute("UPDATE users SET usage_count=GREATEST(0, usage_count-1) WHERE user_id=%s AND bot_name=%s", (uid, BOT_NAME))
        conn.commit()
    except Exception:
        conn.rollback(); raise
    finally:
        conn.close()


def try_use_audit_credit(uid):
    """Атомарне списання одного платного кредиту аудиту."""
    conn = get_db()
    try:
        with conn.cursor() as cur:
            cur.execute('UPDATE users SET audit_credits=audit_credits-1 WHERE user_id=%s AND bot_name=%s AND audit_credits>0', (uid, BOT_NAME))
            ok = cur.rowcount > 0
        conn.commit(); return ok
    except Exception:
        conn.rollback(); raise
    finally:
        conn.close()

def refund_audit_credit(uid):
    """Повернути кредит, якщо доставка не вдалася після списання."""
    conn = get_db()
    try:
        cur = conn.cursor()
        cur.execute('UPDATE users SET audit_credits=audit_credits+1 WHERE user_id=%s AND bot_name=%s', (uid, BOT_NAME))
        conn.commit()
        cur.close()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()

def convert_markdown_tables(text):
    """Перетворення markdown-таблиць на читабельні марковані списки."""
    lines = text.split('\n'); result = []; headers = None
    for line in lines:
        s = line.strip()
        if not s.startswith('|'): headers = None; result.append(line); continue
        cells = [c.strip() for c in s.strip('|').split('|')]
        if all(re.fullmatch(r':?-+:?', c) for c in cells if c): continue
        if headers is None: headers = cells; continue
        parts = [f"{h}: {v}" if h else v for h, v in zip(headers, cells) if v]
        if parts: result.append("• " + " | ".join(parts))
    return '\n'.join(result)

def clean_markdown(text):
    if not text: return text
    text = re.sub(r'^#{1,6}\s*', '', text, flags=re.MULTILINE)
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text); text = re.sub(r'__(.+?)__', r'\1', text)
    text = re.sub(r'\*(.+?)\*', r'\1', text); text = re.sub(r'_(.+?)_', r'\1', text)
    text = re.sub(r'```[a-zA-Z]*\n?', '', text); text = re.sub(r'`(.+?)`', r'\1', text)
    text = re.sub(r'^\s*[\*\-]\s+', '• ', text, flags=re.MULTILINE)
    if '|' in text:
        text = convert_markdown_tables(text)
    return text.strip()

def build_docx(title, body):
    doc = Document(); s = doc.styles['Normal']; s.font.name = 'Calibri'; s.font.size = Pt(11)
    h = doc.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = h.add_run(title.upper()); run.bold = True; run.font.size = Pt(16); run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    doc.add_paragraph()
    for raw in body.split('\n'):
        line = raw.rstrip()
        if not line.strip(): doc.add_paragraph(); continue
        is_h = bool(re.match(r'^\d+\.\s+\S', line))
        ltrs = [c for c in line if c.isalpha()]
        is_caps = len(ltrs) >= 3 and all(c.isupper() for c in ltrs) and len(line) < 80
        if is_h or is_caps:
            p = doc.add_paragraph(); r = p.add_run(line); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        elif line.lstrip().startswith('•'): doc.add_paragraph(line.lstrip()[1:].strip(), style='List Bullet')
        else: doc.add_paragraph(line)
    buf = BytesIO(); doc.save(buf); buf.seek(0); return buf


def safe_filename(n): return re.sub(r'[^\w\-]', '', n.replace(' ', '_'))

def compress_audio(p):
    import subprocess
    output = p + '.compressed.mp3'
    try:
        subprocess.run(['ffmpeg', '-y', '-i', p, '-ac', '1', '-ar', '16000', '-b:a', '32k', output],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=600)
        if os.path.exists(output) and os.path.getsize(output) > 0:
            return output
    except Exception:
        logger.exception('Стиснення аудіо не вдалося; пробуємо оригінал')
    try: os.remove(output)
    except OSError: pass
    return p

def split_audio_chunks(input_path, chunk_minutes=25):
    import subprocess
    out_dir = tempfile.mkdtemp(prefix='skrepka_chunks_')
    pattern = os.path.join(out_dir, 'chunk_%03d.mp3')
    try:
        subprocess.run(['ffmpeg', '-y', '-i', input_path, '-f', 'segment', '-segment_time', str(chunk_minutes * 60),
                        '-ac', '1', '-ar', '16000', '-b:a', '32k', pattern],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=1800)
        chunks = sorted(os.path.join(out_dir, name) for name in os.listdir(out_dir) if name.startswith('chunk_'))
        if not chunks:
            raise RuntimeError('ffmpeg не створив аудіо-фрагментів')
        return chunks
    except Exception:
        shutil.rmtree(out_dir, ignore_errors=True)
        raise

def transcribe_audio_file(path):
    size_mb = os.path.getsize(path) / (1024 * 1024)
    chunk_paths = [path]
    cleanup_files = []
    cleanup_dirs = []
    if size_mb > GROQ_SIZE_LIMIT_MB:
        try:
            chunk_paths = split_audio_chunks(path)
            cleanup_files = list(chunk_paths)
            if chunk_paths:
                cleanup_dirs.append(os.path.dirname(chunk_paths[0]))
            if not chunk_paths:
                raise RuntimeError('ffmpeg не створив аудіо-фрагментів')
        except Exception:
            for d in cleanup_dirs:
                shutil.rmtree(d, ignore_errors=True)
            cleanup_dirs.clear()
            compressed = compress_audio(path)
            chunk_paths = [compressed]
            if compressed != path:
                cleanup_files = [compressed]
    parts = []
    client = _groq_client()
    try:
        for chunk_path in chunk_paths:
            with open(chunk_path, 'rb') as f:
                result = client.audio.transcriptions.create(
                    file=(os.path.basename(chunk_path), f.read()),
                    model='whisper-large-v3-turbo', language='uk')
            text = getattr(result, 'text', '') or ''
            if text.strip():
                parts.append(text.strip())
    finally:
        for item in cleanup_files:
            try: os.remove(item)
            except OSError: pass
        for directory in cleanup_dirs:
            shutil.rmtree(directory, ignore_errors=True)
    return '\n\n'.join(parts)

def generate_document(doc_type, transcription):
    doc_name = DOC_NAMES[doc_type]
    if doc_type == 'transcript': return doc_name, transcription.strip()
    if doc_type not in PROMPTS:
        raise ValueError(f'Невідомий тип документа: {doc_type}')
    resp = _gemini_model().generate_content(f"{PROMPTS[doc_type]}\n\nОсь транскрипція запису:\n\n{transcription}")
    result = getattr(resp, 'text', None)
    if not result or not result.strip():
        raise RuntimeError('AI повернув порожній документ')
    return doc_name, clean_markdown(result)
